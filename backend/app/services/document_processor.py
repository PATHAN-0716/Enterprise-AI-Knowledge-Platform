
from pathlib import Path

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader


def extract_document(file_path: Path) -> dict:
    extension = file_path.suffix.lower()

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document not found.")

    pages = []

    try:
        if extension == ".pdf":
            reader = PdfReader(str(file_path))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Encrypted PDFs are not supported yet.",
                )

            for page_number, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()

                pages.append({
                    "page_number": page_number,
                    "text": text,
                })

        elif extension == ".docx":
            document = Document(str(file_path))

            parts = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    ]
                    parts.append(" | ".join(cells))

            pages.append({
                "page_number": None,
                "text": "\n".join(parts),
            })

        elif extension == ".txt":
            text = file_path.read_text(
                encoding="utf-8",
                errors="replace",
            ).strip()

            pages.append({
                "page_number": 1,
                "text": text,
            })

        else:
            raise HTTPException(
                status_code=400,
                detail="Unsupported document format.",
            )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not extract text from this document.",
        ) from exc

    total_characters = sum(len(page["text"]) for page in pages)

    return {
        "filename": file_path.name,
        "page_count": len(pages),
        "total_characters": total_characters,
        "pages": pages,
    }
